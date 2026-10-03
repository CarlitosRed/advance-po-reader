#!/usr/bin/env python3
"""Build private, review-gated SQLite staging databases from the local sources.

The PDF database retains source files, rendered pages, native page text, and legacy
page-level candidate observations. Candidate targets stay pending human review.
The catalog database imports product data, attributes, and detail-section references;
prices and media assets are excluded by request.

This utility does not call an external model, assign tenants, match products, or train.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import zipfile
from datetime import date, datetime, time, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from PIL import Image
    from pypdf import PdfReader
except ImportError as exc:
    raise SystemExit("Install openpyxl, Pillow, and pypdf in the active Python environment") from exc


CATALOGS = (
    ("wernerco", "WernerCo", "WernerCo-B2B-products.xlsx"),
    ("united_kingdom", "United Kingdom", "UK-B2B-products.xlsx"),
    ("australia", "Australia", "Australia-B2B-products.xlsx"),
)
PRODUCT_SHEET = "Product Information"
ATTRIBUTES_SHEET = "Attributes Values"
DETAILS_SHEET = "Detail Sections"
ASSET_SHEETS = {"Images", "Matrix Image Variation", "External Digital Assets"}
ASSET_COLUMN_MARKERS = ("image", "digital asset")
PRICE_MARKER = "price"
PENDING = "pending_human_review"
PDF_SCHEMA = "pdf-evidence-db-1"
CATALOG_SCHEMA = "regional-product-catalog-db-1"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def raw_text(value: Any) -> str | None:
    """Keep source spelling; a Decimal string is used only for numeric Excel cells."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return Decimal(str(value)).to_eng_string()
    return str(value)


def display_text(cell: Any) -> str | None:
    """Retain simple zero-padded Excel display for numeric catalog identifiers."""
    value = cell.value
    result = raw_text(value)
    if result is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = str(getattr(cell, "number_format", "General") or "General").split(";", 1)[0]
        if re.fullmatch(r"0{2,}", fmt):
            try:
                return str(int(Decimal(str(value)))).zfill(len(fmt))
            except Exception:
                return result
    return result


def cell_json(cell: Any) -> dict[str, Any]:
    return {
        "raw_value": raw_text(cell.value),
        "display_value": display_text(cell),
        "cell_type": str(getattr(cell, "data_type", "")),
        "number_format": str(getattr(cell, "number_format", "General") or "General"),
    }


def rows_with_values(sheet: Any) -> list[tuple[int, tuple[Any, ...]]]:
    result = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        cells = tuple(row)
        if any(cell.value is not None for cell in cells):
            result.append((row_number, cells))
    return result


def private_file(root: Path, relative: str) -> Path:
    root = root.resolve()
    path = (root / relative).resolve()
    if root not in path.parents or not path.is_file():
        raise ValueError("Missing or unsafe path in private VLM snapshot")
    return path


def init_catalog_db(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA foreign_keys=ON;
        CREATE TABLE database_metadata(key TEXT PRIMARY KEY,value_json TEXT NOT NULL);
        CREATE TABLE source_workbooks(
          source_id TEXT PRIMARY KEY,filename TEXT NOT NULL,sha256 TEXT NOT NULL UNIQUE,
          byte_size INTEGER NOT NULL,imported_at_utc TEXT NOT NULL
        );
        CREATE TABLE catalogs(
          catalog_id TEXT PRIMARY KEY,region_label TEXT NOT NULL,
          source_id TEXT NOT NULL REFERENCES source_workbooks(source_id),
          primary_sheet TEXT NOT NULL,
          tenant_assignment_status TEXT NOT NULL
            CHECK(tenant_assignment_status='unassigned_pending_platform_mapping')
        );
        CREATE TABLE source_sheet_inventory(
          catalog_id TEXT NOT NULL REFERENCES catalogs(catalog_id),sheet_name TEXT NOT NULL,
          source_row_count INTEGER NOT NULL,imported_row_count INTEGER NOT NULL,
          included INTEGER NOT NULL CHECK(included IN (0,1)),exclusion_reason TEXT,
          headers_json TEXT NOT NULL,excluded_columns_json TEXT NOT NULL,
          PRIMARY KEY(catalog_id,sheet_name)
        );
        CREATE TABLE catalog_products(
          catalog_id TEXT NOT NULL REFERENCES catalogs(catalog_id),
          product_id_raw TEXT NOT NULL,product_id_display TEXT NOT NULL,
          product_id_cell_type TEXT NOT NULL,product_id_number_format TEXT NOT NULL,
          product_id_source_label TEXT NOT NULL,source_row_number INTEGER NOT NULL,
          product_name_raw TEXT,product_description_raw TEXT,product_features_raw TEXT,
          language_raw TEXT,raw_product_fields_json TEXT NOT NULL,
          PRIMARY KEY(catalog_id,product_id_display),
          UNIQUE(catalog_id,source_row_number)
        );
        CREATE INDEX idx_catalog_products_raw_id
          ON catalog_products(catalog_id,product_id_raw);
        CREATE TABLE product_attributes(
          attribute_id INTEGER PRIMARY KEY,catalog_id TEXT NOT NULL,
          product_id_display TEXT NOT NULL,product_id_raw TEXT NOT NULL,
          source_row_number INTEGER NOT NULL,attribute_label_raw TEXT NOT NULL,
          attribute_value_raw TEXT,attribute_value_display TEXT,
          raw_attribute_fields_json TEXT NOT NULL,
          FOREIGN KEY(catalog_id,product_id_display)
            REFERENCES catalog_products(catalog_id,product_id_display)
        );
        CREATE INDEX idx_product_attributes_product
          ON product_attributes(catalog_id,product_id_display);
        CREATE INDEX idx_product_attributes_label
          ON product_attributes(catalog_id,attribute_label_raw);
        CREATE TABLE product_detail_sections(
          detail_id INTEGER PRIMARY KEY,catalog_id TEXT NOT NULL,
          product_id_display TEXT NOT NULL,product_id_raw TEXT NOT NULL,
          source_row_number INTEGER NOT NULL,detail_section_label_raw TEXT NOT NULL,
          raw_detail_fields_json TEXT NOT NULL,
          FOREIGN KEY(catalog_id,product_id_display)
            REFERENCES catalog_products(catalog_id,product_id_display)
        );
        CREATE INDEX idx_product_detail_sections_product
          ON product_detail_sections(catalog_id,product_id_display);
        CREATE TABLE tenant_catalog_assignments(
          tenant_id TEXT NOT NULL,catalog_id TEXT NOT NULL REFERENCES catalogs(catalog_id),
          assignment_status TEXT NOT NULL CHECK(assignment_status='explicitly_approved'),
          assigned_by TEXT NOT NULL,assigned_at_utc TEXT NOT NULL,
          PRIMARY KEY(tenant_id,catalog_id)
        );
        CREATE TABLE approved_product_aliases(
          alias_id INTEGER PRIMARY KEY,catalog_id TEXT NOT NULL REFERENCES catalogs(catalog_id),
          product_id_display TEXT NOT NULL,alias_raw TEXT NOT NULL,source_label_raw TEXT,
          review_status TEXT NOT NULL CHECK(review_status='human_approved'),
          reviewer TEXT NOT NULL,reason TEXT NOT NULL,
          FOREIGN KEY(catalog_id,product_id_display)
            REFERENCES catalog_products(catalog_id,product_id_display)
        );
        """
    )
    try:
        conn.execute(
            """CREATE VIRTUAL TABLE product_text_search USING fts5(
              catalog_id UNINDEXED,product_id_display UNINDEXED,product_name_raw,
              product_description_raw,product_features_raw,
              tokenize='unicode61 remove_diacritics 2')"""
        )
    except sqlite3.OperationalError as exc:
        if "fts5" not in str(exc).casefold():
            raise
        conn.execute(
            "INSERT INTO database_metadata VALUES(?,?)",
            ("fts5_status", json_text("unavailable_in_sqlite_build")),
        )


def shared_product_ids(conn: sqlite3.Connection) -> int:
    return conn.execute(
        """SELECT COUNT(*) FROM (
             SELECT product_id_display FROM catalog_products
             GROUP BY product_id_display HAVING COUNT(DISTINCT catalog_id)>1
           )"""
    ).fetchone()[0]


def import_catalogs(catalog_dir: Path, output_path: Path) -> dict[str, Any]:
    if output_path.exists():
        raise FileExistsError("Refusing to overwrite an existing catalog database: " + str(output_path))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(output_path.stem + ".building.sqlite3")
    if temp_path.exists():
        raise FileExistsError("Temporary build file already exists: " + str(temp_path))
    conn = sqlite3.connect(temp_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    init_catalog_db(conn)
    result: dict[str, Any] = {"schema_version": CATALOG_SCHEMA, "catalogs": {}}
    try:
        for catalog_id, region, filename in CATALOGS:
            source_path = catalog_dir / filename
            if not source_path.is_file():
                raise FileNotFoundError("Missing source workbook: " + str(source_path))
            source_sha = sha256_file(source_path)
            source_id = "xlsx:" + source_sha
            conn.execute(
                "INSERT INTO source_workbooks VALUES(?,?,?,?,?)",
                (source_id, filename, source_sha, source_path.stat().st_size,
                 datetime.now(timezone.utc).isoformat()),
            )
            book = load_workbook(source_path, read_only=True, data_only=False)
            if PRODUCT_SHEET not in book.sheetnames:
                raise ValueError(filename + " is missing Product Information")
            products = book[PRODUCT_SHEET]
            headers = [str(c.value or "") for c in next(products.iter_rows(min_row=1, max_row=1))]
            if not headers or headers[0] != "Long Product Number":
                raise ValueError(filename + " has an unexpected identifier column")
            if "Product Name" not in headers or "Product Description" not in headers:
                raise ValueError(filename + " is missing product names or descriptions")
            product_rows = rows_with_values(products)
            excluded_columns = [
                h for h in headers
                if PRICE_MARKER in h.casefold()
                or any(marker in h.casefold() for marker in ASSET_COLUMN_MARKERS)
            ]
            included_headers = [h for h in headers if h not in excluded_columns]
            col = {h: i for i, h in enumerate(headers)}
            conn.execute(
                "INSERT INTO catalogs VALUES(?,?,?,?,?)",
                (catalog_id, region, source_id, PRODUCT_SHEET,
                 "unassigned_pending_platform_mapping"),
            )
            counts = {
                "source_file": filename, "source_sha256": source_sha,
                "products": 0, "products_with_name": 0, "products_with_description": 0,
                "attributes": 0, "detail_sections": 0,
                "duplicate_ids": 0, "orphan_attributes": 0, "orphan_detail_sections": 0,
                "excluded_price_columns": [h for h in headers if PRICE_MARKER in h.casefold()],
                "excluded_media_columns": [
                    h for h in headers
                    if any(marker in h.casefold() for marker in ASSET_COLUMN_MARKERS)
                ],
            }
            product_ids: set[str] = set()
            raw_by_display: dict[str, str] = {}
            for row_number, cells in product_rows:
                id_cell = cells[col["Long Product Number"]]
                id_raw, id_display = raw_text(id_cell.value), display_text(id_cell)
                if not id_raw or not id_display:
                    raise ValueError(filename + " has a blank Long Product Number")
                if id_display in product_ids:
                    counts["duplicate_ids"] += 1
                    raise ValueError(filename + " has a duplicate visible product ID")
                product_ids.add(id_display)
                raw_by_display[id_display] = id_raw
                raw_fields = {
                    h: cell_json(cells[i])
                    for i, h in enumerate(headers) if h in included_headers
                }
                name = raw_text(cells[col["Product Name"]].value)
                description = raw_text(cells[col["Product Description"]].value)
                features = raw_text(cells[col["Product Features"]].value) if "Product Features" in col else None
                language = raw_text(cells[col["Language"]].value) if "Language" in col else None
                conn.execute(
                    """INSERT INTO catalog_products(
                         catalog_id,product_id_raw,product_id_display,product_id_cell_type,
                         product_id_number_format,product_id_source_label,source_row_number,
                         product_name_raw,product_description_raw,product_features_raw,
                         language_raw,raw_product_fields_json
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (catalog_id,id_raw,id_display,str(id_cell.data_type),
                     str(id_cell.number_format or "General"),"Long Product Number",row_number,
                     name,description,features,language,json_text(raw_fields)),
                )
                try:
                    conn.execute(
                        """INSERT INTO product_text_search(
                             catalog_id,product_id_display,product_name_raw,
                             product_description_raw,product_features_raw
                           ) VALUES(?,?,?,?,?)""",
                        (catalog_id,id_display,name or "",description or "",features or ""),
                    )
                except sqlite3.OperationalError as exc:
                    if "no such table" not in str(exc).casefold():
                        raise
                counts["products"] += 1
                counts["products_with_name"] += int(bool(name))
                counts["products_with_description"] += int(bool(description))

            for sheet in book.worksheets:
                sheet_headers = [str(c.value or "") for c in next(sheet.iter_rows(min_row=1, max_row=1))]
                sheet_rows = rows_with_values(sheet)
                name = sheet.title
                media_sheet = (
                    name in ASSET_SHEETS
                    or "image" in name.casefold()
                    or "digital asset" in name.casefold()
                )
                include = name in {PRODUCT_SHEET, ATTRIBUTES_SHEET, DETAILS_SHEET}
                if media_sheet:
                    exclusion_reason = "omitted_by_user_scope_media_assets"
                elif include:
                    exclusion_reason = None
                else:
                    exclusion_reason = "not_a_product_or_attribute_or_detail_sheet"
                excluded = [
                    h for h in sheet_headers
                    if PRICE_MARKER in h.casefold()
                    or any(marker in h.casefold() for marker in ASSET_COLUMN_MARKERS)
                ]
                conn.execute(
                    """INSERT INTO source_sheet_inventory VALUES(?,?,?,?,?,?,?,?)""",
                    (catalog_id,name,len(sheet_rows),len(sheet_rows) if include else 0,
                     int(include),exclusion_reason,
                     json_text([h for h in sheet_headers if h not in excluded]),
                     json_text(excluded)),
                )
                if name not in {ATTRIBUTES_SHEET, DETAILS_SHEET} or not sheet_rows:
                    continue
                idx = {h: i for i, h in enumerate(sheet_headers)}
                if "Long Product Number" not in idx:
                    raise ValueError(filename + " contains a child sheet without product IDs")
                for child_row, cells in sheet_rows:
                    id_cell = cells[idx["Long Product Number"]]
                    id_raw, id_display = raw_text(id_cell.value), display_text(id_cell)
                    if id_display not in product_ids:
                        key = "orphan_attributes" if name == ATTRIBUTES_SHEET else "orphan_detail_sections"
                        counts[key] += 1
                        continue
                    raw_fields = {
                        h: cell_json(cells[i]) for i, h in enumerate(sheet_headers)
                        if h not in excluded
                    }
                    if name == ATTRIBUTES_SHEET:
                        label = "Attribute UrlName"
                        value = "Attribute Value Code\n- for all translations -"
                        if label not in idx or value not in idx:
                            raise ValueError(filename + " has unexpected Attributes Values columns")
                        value_cell = cells[idx[value]]
                        conn.execute(
                            """INSERT INTO product_attributes(
                                 catalog_id,product_id_display,product_id_raw,source_row_number,
                                 attribute_label_raw,attribute_value_raw,attribute_value_display,
                                 raw_attribute_fields_json
                               ) VALUES(?,?,?,?,?,?,?,?)""",
                            (catalog_id,id_display,id_raw,child_row,
                             raw_text(cells[idx[label]].value) or "",
                             raw_text(value_cell.value),display_text(value_cell),json_text(raw_fields)),
                        )
                        counts["attributes"] += 1
                    else:
                        label = "Detail Section UrlName"
                        if label not in idx:
                            raise ValueError(filename + " has unexpected Detail Sections columns")
                        conn.execute(
                            """INSERT INTO product_detail_sections(
                                 catalog_id,product_id_display,product_id_raw,source_row_number,
                                 detail_section_label_raw,raw_detail_fields_json
                               ) VALUES(?,?,?,?,?,?)""",
                            (catalog_id,id_display,id_raw,child_row,
                             raw_text(cells[idx[label]].value) or "",json_text(raw_fields)),
                        )
                        counts["detail_sections"] += 1
            book.close()
            if counts["orphan_attributes"] or counts["orphan_detail_sections"]:
                raise ValueError(filename + " has child rows not joined to a source product ID")
            result["catalogs"][catalog_id] = counts

        conn.executemany(
            "INSERT INTO database_metadata VALUES(?,?)",
            [
                ("schema_version",json_text(CATALOG_SCHEMA)),
                ("prices_imported",json_text(False)),
                ("media_assets_imported",json_text(False)),
                ("tenant_assignment",json_text("unassigned; explicit platform mapping required")),
            ],
        )
        conn.execute("PRAGMA user_version=1")
        conn.commit()
        validate_sqlite(conn, "catalog")
        result["totals"] = {
            "products": conn.execute("SELECT COUNT(*) FROM catalog_products").fetchone()[0],
            "attributes": conn.execute("SELECT COUNT(*) FROM product_attributes").fetchone()[0],
            "detail_sections": conn.execute("SELECT COUNT(*) FROM product_detail_sections").fetchone()[0],
            "shared_ids_between_catalogs": shared_product_ids(conn),
            "approved_aliases": conn.execute("SELECT COUNT(*) FROM approved_product_aliases").fetchone()[0],
            "tenant_assignments": conn.execute("SELECT COUNT(*) FROM tenant_catalog_assignments").fetchone()[0],
            "integrity_check": "ok",
        }
        conn.close()
        os.replace(temp_path, output_path)
        return result
    except Exception:
        conn.close()
        if temp_path.exists():
            temp_path.unlink()
        raise


def init_pdf_db(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA foreign_keys=ON;
        CREATE TABLE database_metadata(key TEXT PRIMARY KEY,value_json TEXT NOT NULL);
        CREATE TABLE source_archives(
          archive_id TEXT PRIMARY KEY,filename TEXT NOT NULL,sha256 TEXT NOT NULL UNIQUE,
          byte_size INTEGER NOT NULL,purpose TEXT NOT NULL
        );
        CREATE TABLE source_documents(
          source_sha256 TEXT PRIMARY KEY,corpus_document_id TEXT NOT NULL UNIQUE,
          original_filename TEXT NOT NULL,archive_entry TEXT NOT NULL UNIQUE,
          aliases_json TEXT NOT NULL,media_type TEXT NOT NULL CHECK(media_type='application/pdf'),
          byte_size INTEGER NOT NULL,page_count INTEGER NOT NULL,
          archive_sha256 TEXT NOT NULL REFERENCES source_archives(sha256),
          original_pdf_blob BLOB NOT NULL,leakage_group TEXT,proposed_split_raw TEXT,
          review_status TEXT NOT NULL CHECK(review_status IN
            ('pending_human_review','in_progress','human_approved','rejected')),
          training_eligible INTEGER NOT NULL CHECK(training_eligible IN (0,1))
        );
        CREATE TABLE page_media(
          image_sha256 TEXT PRIMARY KEY,image_rgb_sha256 TEXT NOT NULL,
          width INTEGER NOT NULL,height INTEGER NOT NULL,mime_type TEXT NOT NULL,
          png_blob BLOB NOT NULL
        );
        CREATE TABLE document_pages(
          source_sha256 TEXT NOT NULL REFERENCES source_documents(source_sha256),
          page_number INTEGER NOT NULL,page_id TEXT NOT NULL,
          source_page_sha256 TEXT NOT NULL,source_page_hash_method TEXT NOT NULL,
          image_sha256 TEXT NOT NULL REFERENCES page_media(image_sha256),
          image_rgb_sha256 TEXT NOT NULL,render_method TEXT NOT NULL,dpi INTEGER NOT NULL,
          width INTEGER NOT NULL,height INTEGER NOT NULL,native_text_raw TEXT NOT NULL,
          native_text_sha256 TEXT NOT NULL,native_text_chars INTEGER NOT NULL,
          candidate_target_json TEXT NOT NULL,candidate_sha256 TEXT NOT NULL,
          candidate_file_sha256 TEXT NOT NULL,annotation_origin TEXT NOT NULL,
          layout_family TEXT,leakage_group TEXT,proposed_split_raw TEXT,page_status TEXT NOT NULL,
          review_status TEXT NOT NULL CHECK(review_status IN
            ('pending_human_review','in_progress','human_approved','rejected')),
          training_eligible INTEGER NOT NULL CHECK(training_eligible IN (0,1)),
          PRIMARY KEY(source_sha256,page_number)
        );
        CREATE INDEX idx_document_pages_image ON document_pages(image_rgb_sha256);
        CREATE TABLE line_observations(
          observation_id TEXT PRIMARY KEY,source_sha256 TEXT NOT NULL,page_number INTEGER NOT NULL,
          line_number INTEGER NOT NULL,page_image_rgb_sha256 TEXT NOT NULL,visual_line_key TEXT NOT NULL,
          visual_duplicate_occurrence_count INTEGER NOT NULL DEFAULT 1,
          primary_visual_occurrence INTEGER NOT NULL DEFAULT 1 CHECK(primary_visual_occurrence IN (0,1)),
          kind_raw TEXT,source_line_label_raw TEXT,identifiers_json TEXT NOT NULL,
          description_raw TEXT,quantity_raw TEXT,uom_raw TEXT,amount_raw TEXT,unit_price_raw TEXT,
          attributes_json TEXT NOT NULL,evidence_text_raw TEXT,uncertainties_json TEXT NOT NULL,
          raw_candidate_line_json TEXT NOT NULL,
          review_status TEXT NOT NULL CHECK(review_status IN
            ('pending_human_review','in_progress','human_approved','rejected')),
          training_eligible INTEGER NOT NULL CHECK(training_eligible IN (0,1)),
          FOREIGN KEY(source_sha256,page_number) REFERENCES document_pages(source_sha256,page_number),
          UNIQUE(source_sha256,page_number,line_number)
        );
        CREATE INDEX idx_line_observations_page ON line_observations(source_sha256,page_number,line_number);
        CREATE TABLE human_review_events(
          event_id TEXT PRIMARY KEY,source_sha256 TEXT NOT NULL REFERENCES source_documents(source_sha256),
          page_number INTEGER,reviewer_name TEXT NOT NULL CHECK(length(trim(reviewer_name))>0),
          decision TEXT NOT NULL CHECK(decision IN ('approve','reject','correction','note')),
          event_utc TEXT NOT NULL,reason TEXT NOT NULL,previous_values_json TEXT,
          corrected_values_json TEXT,evidence_json TEXT NOT NULL,
          FOREIGN KEY(source_sha256,page_number) REFERENCES document_pages(source_sha256,page_number)
        );
        CREATE TRIGGER human_review_events_no_update BEFORE UPDATE ON human_review_events
          BEGIN SELECT RAISE(ABORT,'Human review events are append-only'); END;
        CREATE TRIGGER human_review_events_no_delete BEFORE DELETE ON human_review_events
          BEGIN SELECT RAISE(ABORT,'Human review events are append-only'); END;
        CREATE VIEW unique_visual_line_observations AS
          SELECT * FROM line_observations WHERE primary_visual_occurrence=1;
        CREATE VIEW approved_training_examples AS
          SELECT d.corpus_document_id,d.source_sha256,p.page_number,p.image_sha256,p.candidate_target_json
          FROM source_documents d JOIN document_pages p ON p.source_sha256=d.source_sha256
          WHERE d.review_status='human_approved' AND d.training_eligible=1
            AND p.review_status='human_approved' AND p.training_eligible=1
            AND EXISTS(SELECT 1 FROM human_review_events r
              WHERE r.source_sha256=d.source_sha256 AND r.page_number=p.page_number
                AND r.decision='approve' AND length(trim(r.reviewer_name))>0)
            AND NOT EXISTS(SELECT 1 FROM document_pages pending
              WHERE pending.source_sha256=d.source_sha256
                AND (pending.review_status<>'human_approved' OR pending.training_eligible<>1
                  OR NOT EXISTS(SELECT 1 FROM human_review_events reviewed
                    WHERE reviewed.source_sha256=pending.source_sha256
                      AND reviewed.page_number=pending.page_number
                      AND reviewed.decision='approve'
                      AND length(trim(reviewed.reviewer_name))>0)));
        """
    )
    try:
        conn.execute(
            """CREATE VIRTUAL TABLE page_text_search USING fts5(
              corpus_document_id UNINDEXED,page_number UNINDEXED,native_text_raw,
              tokenize='unicode61 remove_diacritics 2')"""
        )
    except sqlite3.OperationalError as exc:
        if "fts5" not in str(exc).casefold():
            raise
        conn.execute(
            "INSERT INTO database_metadata VALUES(?,?)",
            ("fts5_status",json_text("unavailable_in_sqlite_build")),
        )


def validate_sqlite(conn: sqlite3.Connection, label: str) -> None:
    integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
    fk_errors = list(conn.execute("PRAGMA foreign_key_check"))
    if integrity != "ok" or fk_errors:
        raise RuntimeError(label + " database validation failed: integrity=" + str(integrity))


def pixel_hash(image: Any) -> str:
    rgb = image.convert("RGB")
    width, height = rgb.size
    prefix = ("RGB:" + str(width) + ":" + str(height) + ":").encode("ascii")
    return sha256_bytes(prefix + rgb.tobytes())


def import_pdfs(source_zip: Path, page_dataset: Path, output_path: Path) -> dict[str, Any]:
    if output_path.exists():
        raise FileExistsError("Refusing to overwrite an existing PDF database: " + str(output_path))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(output_path.stem + ".building.sqlite3")
    if temp_path.exists():
        raise FileExistsError("Temporary build file already exists: " + str(temp_path))
    manifest = json.loads((page_dataset / "manifest.json").read_text(encoding="utf-8"))
    archive_bytes = source_zip.read_bytes()
    archive_sha = sha256_bytes(archive_bytes)
    if archive_sha != manifest.get("archive", {}).get("sha256"):
        raise ValueError("Source ZIP hash does not match the page dataset")
    conn = sqlite3.connect(temp_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    init_pdf_db(conn)
    summary = {
        "schema_version": PDF_SCHEMA,"archive_sha256": archive_sha,
        "documents": 0,"pages": 0,"unique_page_images": 0,
        "candidate_lines": 0,"unique_visual_candidate_lines": 0,
        "duplicate_extra_visual_lines": 0,"pending_pages": 0,
        "approved_training_examples": 0,"documents_by_id": {},
    }
    try:
        with zipfile.ZipFile(source_zip, "r") as archive:
            pdf_entries = {
                item.filename for item in archive.infolist()
                if item.filename.casefold().endswith(".pdf") and not item.is_dir()
            }
            manifest_entries = {
                alias["archive_entry"]
                for doc in manifest["documents"] for alias in doc.get("aliases", [])
            }
            if len(pdf_entries) != len(manifest["documents"]) or pdf_entries != manifest_entries:
                raise ValueError("PDF ZIP entries do not match the complete page-dataset manifest")
            conn.execute(
                "INSERT INTO source_archives VALUES(?,?,?,?,?)",
                ("archive:" + archive_sha,source_zip.name,archive_sha,len(archive_bytes),
                 "private supplied PDF corpus; local evidence and human review"),
            )
            for doc in manifest["documents"]:
                sha = doc["source_sha256"]
                aliases = doc.get("aliases", [])
                if not aliases:
                    raise ValueError("A source PDF has no archive alias")
                alias = aliases[0]
                entry = alias["archive_entry"]
                pdf = archive.read(entry)
                if sha256_bytes(pdf) != sha or len(pdf) != int(doc["size_bytes"]):
                    raise ValueError("PDF bytes or size do not match the source manifest")
                if len(PdfReader(BytesIO(pdf)).pages) != len(doc["pages"]):
                    raise ValueError("PDF page count does not match the page evidence")
                conn.execute(
                    """INSERT INTO source_documents(
                         source_sha256,corpus_document_id,original_filename,archive_entry,aliases_json,
                         media_type,byte_size,page_count,archive_sha256,original_pdf_blob,
                         leakage_group,proposed_split_raw,review_status,training_eligible
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (sha,doc["document_id"],alias["filename"],entry,json_text(aliases),
                     "application/pdf",len(pdf),len(doc["pages"]),archive_sha,
                     sqlite3.Binary(pdf),doc.get("leakage_group"),doc.get("proposed_split"),PENDING,0),
                )
                summary["documents"] += 1
                summary["documents_by_id"][doc["document_id"]] = {
                    "filename":alias["filename"],"sha256":sha,"pages":len(doc["pages"]),
                    "review_status":PENDING,"training_eligible":False,
                }
                for page in doc["pages"]:
                    image_path = private_file(page_dataset,page["image_path"])
                    image_bytes = image_path.read_bytes()
                    if sha256_bytes(image_bytes) != page["image_sha256"]:
                        raise ValueError("Rendered page file hash does not match the manifest")
                    with Image.open(image_path) as image:
                        image.load()
                        width,height = image.size
                        if pixel_hash(image) != page["image_rgb_sha256"]:
                            raise ValueError("Rendered page pixels do not match the manifest")
                    if width != int(page["width"]) or height != int(page["height"]):
                        raise ValueError("Rendered page dimensions do not match the manifest")
                    if not conn.execute("SELECT 1 FROM page_media WHERE image_sha256=?",(page["image_sha256"],)).fetchone():
                        conn.execute(
                            "INSERT INTO page_media VALUES(?,?,?,?,?,?)",
                            (page["image_sha256"],page["image_rgb_sha256"],width,height,
                             "image/png",sqlite3.Binary(image_bytes)),
                        )
                    native_path = private_file(page_dataset,page["native_text_path"])
                    native_bytes = native_path.read_bytes()
                    if sha256_bytes(native_bytes) != page["native_text_sha256"]:
                        raise ValueError("Native page text hash does not match the manifest")
                    native_text = native_bytes.decode("utf-8")
                    candidate_path = private_file(page_dataset,page["candidate_path"])
                    candidate_bytes = candidate_path.read_bytes()
                    if sha256_bytes(candidate_bytes) != page["candidate_file_sha256"]:
                        raise ValueError("Page candidate hash does not match the manifest")
                    candidate = json.loads(candidate_bytes.decode("utf-8"))
                    if (candidate.get("source_sha256") != sha
                        or int(candidate.get("page",0)) != int(page["page"])
                        or candidate.get("image_sha256") != page["image_sha256"]):
                        raise ValueError("Page target is not bound to the matching source page")
                    if candidate.get("review_status") != PENDING or candidate.get("page_status") != "candidate":
                        raise ValueError("Only pending page candidates can enter this staging database")
                    if not isinstance(candidate.get("target"),dict):
                        raise ValueError("Page target must be a structured legacy observation object")
                    conn.execute(
                        """INSERT INTO document_pages(
                           source_sha256,page_number,page_id,source_page_sha256,source_page_hash_method,
                           image_sha256,image_rgb_sha256,render_method,dpi,width,height,native_text_raw,
                           native_text_sha256,native_text_chars,candidate_target_json,candidate_sha256,
                           candidate_file_sha256,annotation_origin,layout_family,leakage_group,
                           proposed_split_raw,page_status,review_status,training_eligible
                         ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (sha,int(page["page"]),page["page_id"],page["source_page_sha256"],
                         page["source_page_hash_method"],page["image_sha256"],page["image_rgb_sha256"],
                         page["render_method"],int(page["dpi"]),width,height,native_text,
                         page["native_text_sha256"],len(native_text),json_text(candidate),
                         page["candidate_sha256"],page["candidate_file_sha256"],
                         candidate["annotation_origin"],page.get("layout_family"),
                         page.get("leakage_group"),page.get("proposed_split"),
                         candidate["page_status"],PENDING,0),
                    )
                    try:
                        conn.execute(
                            "INSERT INTO page_text_search VALUES(?,?,?)",
                            (doc["document_id"],int(page["page"]),native_text),
                        )
                    except sqlite3.OperationalError as exc:
                        if "no such table" not in str(exc).casefold():
                            raise
                    for number,line in enumerate(candidate["target"].get("lines",[]),start=1):
                        if not isinstance(line,dict):
                            raise ValueError("Target line is not an object")
                        conn.execute(
                            """INSERT INTO line_observations(
                               observation_id,source_sha256,page_number,line_number,
                               page_image_rgb_sha256,visual_line_key,kind_raw,source_line_label_raw,
                               identifiers_json,description_raw,quantity_raw,uom_raw,amount_raw,
                               unit_price_raw,attributes_json,evidence_text_raw,uncertainties_json,
                               raw_candidate_line_json,review_status,training_eligible
                             ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                            (sha[:16]+"-p"+str(page["page"]).zfill(3)+"-l"+str(number).zfill(3),
                             sha,int(page["page"]),number,page["image_rgb_sha256"],
                             page["image_rgb_sha256"]+":line:"+str(number),
                             raw_text(line.get("kind")),raw_text(line.get("source_line_label")),
                             json_text(line.get("identifiers",[])),raw_text(line.get("description_raw")),
                             raw_text(line.get("quantity_raw")),raw_text(line.get("uom_raw")),
                             raw_text(line.get("amount_raw")),raw_text(line.get("unit_price_raw")),
                             json_text(line.get("attributes",[])),raw_text(line.get("evidence_text")),
                             json_text(line.get("uncertainties",[])),json_text(line),PENDING,0),
                        )
                        summary["candidate_lines"] += 1
                    summary["pages"] += 1
                    summary["pending_pages"] += 1
        groups = conn.execute(
            """SELECT page_image_rgb_sha256,line_number,COUNT(*) AS n
               FROM line_observations GROUP BY page_image_rgb_sha256,line_number"""
        ).fetchall()
        for group in groups:
            rows = conn.execute(
                """SELECT observation_id FROM line_observations
                   WHERE page_image_rgb_sha256=? AND line_number=?
                   ORDER BY source_sha256,page_number,observation_id""",
                (group["page_image_rgb_sha256"],group["line_number"]),
            ).fetchall()
            for rank,row in enumerate(rows):
                conn.execute(
                    """UPDATE line_observations
                       SET visual_duplicate_occurrence_count=?,primary_visual_occurrence=?
                       WHERE observation_id=?""",
                    (group["n"],int(rank==0),row["observation_id"]),
                )
        summary["unique_page_images"] = conn.execute("SELECT COUNT(*) FROM page_media").fetchone()[0]
        summary["unique_visual_candidate_lines"] = conn.execute(
            "SELECT COUNT(*) FROM unique_visual_line_observations"
        ).fetchone()[0]
        summary["duplicate_extra_visual_lines"] = conn.execute(
            """SELECT COALESCE(SUM(n-1),0) FROM (
                 SELECT COUNT(*) AS n FROM line_observations
                 GROUP BY page_image_rgb_sha256,line_number HAVING COUNT(*)>1
               )"""
        ).fetchone()[0]
        conn.executemany(
            "INSERT INTO database_metadata VALUES(?,?)",
            [
                ("schema_version",json_text(PDF_SCHEMA)),
                ("archive_sha256",json_text(archive_sha)),
                ("legacy_target_schema",json_text(manifest["target_schema_version"])),
                ("product_document_v1_status",json_text("adapter pending; legacy page observations retained")),
                ("training_status",json_text("blocked_pending_human_review")),
                ("review_scope",json_text("all pages in a source document require named human review")),
                ("external_provider_calls",json_text(0)),
            ],
        )
        conn.execute("PRAGMA user_version=1")
        conn.commit()
        validate_sqlite(conn,"PDF evidence")
        summary["approved_training_examples"] = conn.execute(
            "SELECT COUNT(*) FROM approved_training_examples"
        ).fetchone()[0]
        conn.close()
        os.replace(temp_path,output_path)
        return summary
    except Exception:
        conn.close()
        if temp_path.exists():
            temp_path.unlink()
        raise


def catalog_totals(path: Path) -> dict[str, Any]:
    with sqlite3.connect(path) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            """SELECT c.catalog_id,c.region_label,COUNT(p.product_id_display) products,
                      SUM(CASE WHEN COALESCE(p.product_name_raw,'')<>'' THEN 1 ELSE 0 END) named,
                      SUM(CASE WHEN COALESCE(p.product_description_raw,'')<>'' THEN 1 ELSE 0 END) described,
                      (SELECT COUNT(*) FROM product_attributes a WHERE a.catalog_id=c.catalog_id) attributes,
                      (SELECT COUNT(*) FROM product_detail_sections s WHERE s.catalog_id=c.catalog_id) details
               FROM catalogs c JOIN catalog_products p ON p.catalog_id=c.catalog_id
               GROUP BY c.catalog_id,c.region_label ORDER BY c.catalog_id"""
        ).fetchall()
        return {
            row["catalog_id"]:{
                "region_label":row["region_label"],"products":row["products"],
                "products_with_name":row["named"],"products_with_description":row["described"],
                "attribute_rows":row["attributes"],"detail_section_rows":row["details"],
            } for row in rows
        }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-zip",type=Path,default=Path("data/source-snapshot/corpus/full-data-set.zip"))
    parser.add_argument("--page-dataset",type=Path,default=Path("output/source-databases/vlm-dataset"))
    parser.add_argument("--catalog-dir",type=Path,default=Path("data/source-snapshot/catalogs"))
    parser.add_argument("--output-dir",type=Path,default=Path("output/source-databases"))
    args = parser.parse_args(argv)
    args.output_dir.mkdir(parents=True,exist_ok=True)
    pdf_path = args.output_dir/"pdf-evidence.sqlite3"
    catalog_path = args.output_dir/"regional-product-catalog.sqlite3"
    report_path = args.output_dir/"source-databases-manifest.json"
    if report_path.exists():
        parser.error("Refusing to replace an existing manifest; build into a new output directory")
    pdf = import_pdfs(args.source_zip,args.page_dataset,pdf_path)
    catalog = import_catalogs(args.catalog_dir,catalog_path)
    report = {
        "status":"private_local_staging_databases",
        "database_files":{
            pdf_path.name:{"sha256":sha256_file(pdf_path),"size_bytes":pdf_path.stat().st_size},
            catalog_path.name:{"sha256":sha256_file(catalog_path),"size_bytes":catalog_path.stat().st_size},
        },
        "pdf_evidence":pdf,"product_catalogs":catalog,
        "catalog_totals":catalog_totals(catalog_path),
        "training":{
            "approved_examples":pdf["approved_training_examples"],
            "candidate_pages_pending_human_review":pdf["pending_pages"],
            "eligible_for_training_now":False,
            "reason":"The source page targets are pending visual transcription candidates; no complete human review has been recorded.",
        },
        "platform_integration":{
            "format":"SQLite private local staging; no FastAPI/PostgreSQL connection was implemented or claimed",
            "tenant_catalog_assignments":0,
            "product_document_v1":"adapter pending; targets retain the legacy page-observation contract",
            "external_provider_calls":0,
        },
    }
    report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({
        "manifest":str(report_path),"pdf_database":str(pdf_path),"catalog_database":str(catalog_path),
        "pdfs":pdf["documents"],"pages":pdf["pages"],"candidate_lines":pdf["candidate_lines"],
        "unique_visual_lines":pdf["unique_visual_candidate_lines"],
        "approved_training_examples":pdf["approved_training_examples"],
        "catalog_totals":report["catalog_totals"],
    },ensure_ascii=False,indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
